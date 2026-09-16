"""独立 document service 生成资源后，agent service 读取并执行问答的集成测试。"""

import io
import json
import uuid

import grpc
import numpy as np
import pytest
from docx import Document

from agent_proto import agent_pb2 as pb
from document_service.document_resources import model as embedding_model
from traceagent_shared.object_store import parse_resource_path


@pytest.fixture
def session_id():
    """每个测试用独立会话 id，资源发布进各自的 res_<session_id> 桶。"""
    return uuid.uuid4().hex[:12]


@pytest.fixture
def resources(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCUMENT_RESOURCES_ROOT", str(tmp_path))
    calls = []

    class Embedder:
        def tokenize(self, text):
            return [(i, i + 1) for i in range(len(text))]

        def encode(self, texts):
            calls.extend(texts)
            return np.array([[1.0, 0.0] for _ in texts], dtype=np.float32)

    monkeypatch.setattr(embedding_model, "get_embedder", lambda **kwargs: Embedder())
    return tmp_path, calls


def upload(document_rpc, session_id):
    document = Document()
    document.add_heading("合同", 1)
    document.add_paragraph("付款期限为三十天。")
    data = io.BytesIO()
    document.save(data)
    return document_rpc.PrepareResources(pb.PrepareResourcesRequest(session_id=session_id, files=[
        pb.UploadedFile(filename="合同.docx", content=data.getvalue()),
        pb.UploadedFile(filename="附件.docx", content=data.getvalue()),
    ]), timeout=10)


def chat_request(refs, completion_id="cmp_resource"):
    return pb.ChatCompletionRequest(
        completion_id=completion_id,
        resource_path=[pb.ResourceRef(type=ref.type, location=ref.location) for ref in refs],
        messages=[pb.QaMessage(role="user", content="你好")],
    )


def _bucket(refs):
    documents_location = next(ref.location for ref in refs if ref.type == "documents")
    bucket, _ = parse_resource_path(documents_location)
    return bucket


def test_qa_uses_prepared_path_without_rebuilding_or_deleting(resources, document_rpc, rpc, monkeypatch, s3_store, session_id):
    """document service 生成的资源可被 agent 两轮复用，问答不会重建或删除对象。"""
    from tests.file_extraction_agent.test_streaming_retry import StreamingModel
    from service.file_extraction_agent import application as qa_route

    model = StreamingModel()
    model._release.set()
    monkeypatch.setattr(qa_route, "build_qa_model", lambda config: model)
    refs = list(upload(document_rpc, session_id).resource_path)
    bucket = _bucket(refs)
    before = sorted(s3_store.list_objects(bucket))
    for cid in ("cmp_first", "cmp_second"):
        events = list(rpc.ChatCompletion(chat_request(refs, completion_id=cid), timeout=5))
        assert events[-1].type == "completion.completed"
        assert "".join(e.delta for e in events if e.type == "model_message.delta") == "前半后半"
        assert sum(e.type == "model_message.done" for e in events) == 1
    assert sorted(s3_store.list_objects(bucket)) == before


def test_qa_rejects_unmanaged_resource_path(resources, rpc):
    """缺少 documents/index 定位的资源引用在首事件前返回 INVALID_ARGUMENT。"""
    with pytest.raises(grpc.RpcError) as error:
        next(rpc.ChatCompletion(pb.ChatCompletionRequest(
            completion_id="cmp_bad",
            resource_path=[pb.ResourceRef(type="raw", location="s3://nonexistent/raw/a.pdf")],
            messages=[pb.QaMessage(role="user", content="你好")],
        ), timeout=5))
    assert error.value.code() == grpc.StatusCode.INVALID_ARGUMENT


@pytest.mark.parametrize("damage", ["missing_index", "bad_version", "outside_reference"])
def test_qa_rejects_damaged_resource_without_rebuilding(resources, document_rpc, rpc, s3_store, damage, session_id):
    """索引缺失、清单版本错误、引用越界均由 agent 拒绝且不重建。"""
    refs = list(upload(document_rpc, session_id).resource_path)
    bucket = _bucket(refs)
    if damage == "missing_index":
        s3_store.delete_object(bucket, "index/vectors.npy")
    elif damage == "bad_version":
        manifest = json.loads(s3_store.get_object(bucket, "manifest.json").decode("utf-8"))
        manifest["version"] = -1
        s3_store.put_object(bucket, "manifest.json", json.dumps(manifest).encode("utf-8"))
    else:
        index = json.loads(s3_store.get_object(bucket, "index/index.json").decode("utf-8"))
        index["chunks"][0]["covered_files"] = ["../../outside.md"]
        s3_store.put_object(bucket, "index/index.json", json.dumps(index).encode("utf-8"))
    with pytest.raises(grpc.RpcError) as error:
        next(rpc.ChatCompletion(chat_request(refs), timeout=5))
    assert error.value.code() == grpc.StatusCode.INVALID_ARGUMENT
