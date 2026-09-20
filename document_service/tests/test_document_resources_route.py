"""真实 DOCX → gRPC 上传 → HTML/资源发布（storage 服务）→ 路径问答，验证完整资源链路。"""

import io
import json
import uuid
import zipfile
from pathlib import Path

import grpc
import numpy as np
import pytest
from docx import Document

from agent_proto import agent_pb2 as pb
from document_service.document_resources import model as embedding_model


@pytest.fixture
def resources(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCUMENT_RESOURCES_ROOT", str(tmp_path))
    calls = []
    class Embedder:
        def tokenize(self, text):
            return [(i, i + 1) for i in range(len(text))]

        def encode(self, texts):
            calls.extend(texts)
            return np.array([[1.0, 0.0] for _ in texts], dtype=np.float32)
    monkeypatch.setattr(embedding_model, "get_embedder", lambda **kwargs: Embedder())
    return tmp_path, calls


@pytest.fixture
def session_id():
    """每个测试用独立会话 id，避免共享 storage 服务里同名桶互相污染。"""
    return uuid.uuid4().hex[:12]


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("合同", 1)
    document.add_paragraph("付款期限为三十天。")
    data = io.BytesIO()
    document.save(data)
    return data.getvalue()


def prepare(rpc, session_id, *, files=None, remove_raw=()):
    """发送一次 PrepareResources；缺省上传两个 DOCX，可传空批次加 remove_raw。"""
    if files is None:
        files = [pb.UploadedFile(filename="合同.docx", content=_docx_bytes()),
                 pb.UploadedFile(filename="附件.docx", content=_docx_bytes())]
    return rpc.PrepareResources(pb.PrepareResourcesRequest(
        session_id=session_id, files=files, remove_raw=remove_raw), timeout=10)


def _bucket(refs):
    from traceagent_shared.object_store import parse_resource_path
    documents_location = next(ref.location for ref in refs if ref.type == "documents")
    bucket, _ = parse_resource_path(documents_location)
    return bucket


def _raw_location(refs, filename):
    return next(ref.location for ref in refs
                if ref.type == "raw" and ref.location.endswith(f"/raw/{filename}"))


def test_prepare_real_docx_publishes_complete_resource(resources, document_rpc, s3_store, session_id):
    """真实多文档上传发布单个文档树归档和可用索引，桶固定为 res_<session_id>。"""
    root, calls = resources
    result = prepare(document_rpc, session_id)
    refs = list(result.resource_path)
    assert [ref.type for ref in refs] == ["documents", "index", "raw", "raw"]
    bucket = _bucket(refs)
    assert bucket == f"res_{session_id}"
    assert s3_store.get_object(bucket, "documents.zip") is not None
    assert s3_store.get_object(bucket, "manifest.json") is not None
    assert s3_store.get_object(bucket, "index/vectors.npy") is not None
    assert s3_store.get_object(bucket, "index/index.json") is not None
    assert s3_store.list_objects(bucket, prefix="documents/") == []  # 树只存在于归档里
    assert calls


def test_second_upload_reuses_session_bucket_and_keeps_previous_raws(resources, document_rpc, s3_store, session_id):
    """同一会话第二次上传仍写入同一桶，返回的 raw 引用覆盖旧文件和新文件。"""
    prepare(document_rpc, session_id, files=[pb.UploadedFile(filename="a.docx", content=_docx_bytes())])
    second = prepare(document_rpc, session_id, files=[pb.UploadedFile(filename="b.docx", content=_docx_bytes())])
    refs = list(second.resource_path)
    bucket = _bucket(refs)
    assert bucket == f"res_{session_id}"
    assert [ref.location.rsplit("/", 1)[1] for ref in refs if ref.type == "raw"] == ["a.docx", "b.docx"]
    assert s3_store.get_object(bucket, "raw/a.docx") is not None
    assert s3_store.get_object(bucket, "raw/b.docx") is not None


def test_remove_raw_with_empty_files_deletes_object_and_rebuilds(resources, document_rpc, s3_store, session_id):
    """backend remove_file 的调用形状：空批次 + remove_raw，删除桶内 raw 并用剩余文件重建。"""
    first = prepare(document_rpc, session_id)
    refs = list(first.resource_path)
    target = pb.ResourceRef(type="raw", location=_raw_location(refs, "合同.docx"))
    result = prepare(document_rpc, session_id, files=[], remove_raw=[target])
    remaining = list(result.resource_path)
    bucket = _bucket(refs)
    assert [ref.location.rsplit("/", 1)[1] for ref in remaining if ref.type == "raw"] == ["附件.docx"]
    assert s3_store.list_objects(bucket, prefix="raw/") == ["raw/附件.docx"]
    assert s3_store.get_object(bucket, "documents.zip") is not None
    assert s3_store.get_object(bucket, "index/index.json") is not None


def test_remove_last_raw_clears_published_artifacts(resources, document_rpc, s3_store, session_id):
    """移除会话最后一个 raw 后返回空引用，桶内不再残留归档和索引。"""
    first = prepare(document_rpc, session_id, files=[pb.UploadedFile(filename="a.docx", content=_docx_bytes())])
    refs = list(first.resource_path)
    bucket = _bucket(refs)
    target = pb.ResourceRef(type="raw", location=_raw_location(refs, "a.docx"))
    result = prepare(document_rpc, session_id, files=[], remove_raw=[target])
    assert list(result.resource_path) == []
    assert s3_store.list_objects(bucket) == []


def test_remove_raw_rejects_other_bucket_or_type(resources, document_rpc, session_id):
    """remove_raw 的 type 必须为 raw，且引用必须属于本会话桶，否则 INVALID_ARGUMENT。"""
    for ref in (pb.ResourceRef(type="documents", location=f"s3://res_{session_id}/documents.zip"),
                pb.ResourceRef(type="raw", location="s3://res_other/raw/a.pdf"),
                pb.ResourceRef(type="raw", location="not-s3")):
        with pytest.raises(grpc.RpcError) as error:
            prepare(document_rpc, session_id, files=[], remove_raw=[ref])
        assert error.value.code() == grpc.StatusCode.INVALID_ARGUMENT


def test_prepare_requires_session_id(resources, document_rpc):
    """缺少 session_id 无法确定会话桶，返回 INVALID_ARGUMENT。"""
    with pytest.raises(grpc.RpcError) as error:
        prepare(document_rpc, "")
    assert error.value.code() == grpc.StatusCode.INVALID_ARGUMENT


def test_protocol_prepare_response_drops_documents_payload():
    """共享协议不再暴露 documents/Document，调用方只拿资源定位。"""
    fields = pb.PrepareResourcesResponse.DESCRIPTOR.fields_by_name
    assert "resource_path" in fields
    assert "documents" not in fields
    assert "Document" not in pb.DESCRIPTOR.message_types_by_name


def test_read_blocks_returns_block_text_from_archive(resources, document_rpc, s3_store, session_id):
    """归档内 key 返回原文；不存在的 key 返回 found=false。"""
    root, calls = resources
    refs = list(prepare(document_rpc, session_id).resource_path)
    bucket = _bucket(refs)
    archive = s3_store.get_object(bucket, "documents.zip")
    members = zipfile.ZipFile(io.BytesIO(archive)).namelist()
    key = next(name for name in members if name.endswith(".md"))
    result = document_rpc.ReadBlocks(pb.ReadBlocksRequest(bucket=bucket, keys=[key, "documents/missing.md"]), timeout=5)
    blocks = list(result.blocks)
    assert [block.key for block in blocks] == [key, "documents/missing.md"]
    assert "付款期限为三十天。" in blocks[0].text
    assert blocks[0].found
    assert not blocks[1].found
    assert blocks[1].text == ""


def test_read_blocks_requires_bucket_and_keys(resources, document_rpc):
    """空 bucket 或空 keys 返回 INVALID_ARGUMENT。"""
    for request in (pb.ReadBlocksRequest(bucket="", keys=["documents/a.md"]),
                    pb.ReadBlocksRequest(bucket="res_x", keys=[])):
        with pytest.raises(grpc.RpcError) as error:
            document_rpc.ReadBlocks(request, timeout=5)
        assert error.value.code() == grpc.StatusCode.INVALID_ARGUMENT


def test_read_blocks_missing_archive_is_not_found(resources, document_rpc):
    """归档缺失（未上传会话）返回 NOT_FOUND。"""
    with pytest.raises(grpc.RpcError) as error:
        document_rpc.ReadBlocks(pb.ReadBlocksRequest(bucket="res_missing", keys=["documents/a.md"]), timeout=5)
    assert error.value.code() == grpc.StatusCode.NOT_FOUND


def test_prepare_failure_does_not_publish_resource(resources, document_rpc, monkeypatch, session_id):
    """embedding 失败映射 INTERNAL，不发布半成品。"""
    def fail(**kwargs):
        raise RuntimeError("embedding unavailable")
    monkeypatch.setattr(embedding_model, "get_embedder", fail)
    with pytest.raises(grpc.RpcError) as error:
        prepare(document_rpc, session_id)
    assert error.value.code() == grpc.StatusCode.INTERNAL
    assert "embedding unavailable" in error.value.details()


@pytest.mark.parametrize("files", [[], [pb.UploadedFile(filename="bad.txt", content=b"text")]])
def test_prepare_rejects_unsupported_or_missing_files(resources, document_rpc, session_id, files):
    """空上传或不支持的类型在处理前返回 INVALID_ARGUMENT。"""
    with pytest.raises(grpc.RpcError) as error:
        prepare(document_rpc, session_id, files=files)
    assert error.value.code() == grpc.StatusCode.INVALID_ARGUMENT


def test_prepare_pdf_calls_parser_then_builds_resource(resources, document_rpc, monkeypatch, session_id):
    """PDF bytes 交给解析器，返回的 HTML 进入真实资源构建。"""
    from document_service.document_processor import processor
    from document_service.document_processor.schemas import ProcessResult
    calls = []
    def parse(file, file_type=None):
        calls.append((file.filename, file.read()))
        return ProcessResult(filename=file.filename, html="<p>PDF 内容</p>")
    monkeypatch.setattr(processor, "process", parse)
    result = prepare(document_rpc, session_id, files=[pb.UploadedFile(filename="a.pdf", content=b"%PDF-1.4")])
    assert calls == [("a.pdf", b"%PDF-1.4")]
    assert _bucket(list(result.resource_path)) == f"res_{session_id}"
    assert resources[1] == ["PDF 内容"]


def test_parser_failure_identifies_file_and_does_not_build_index(resources, document_rpc, monkeypatch, session_id):
    """解析异常映射 INTERNAL 并标明文件，不进入索引构建。"""
    from document_service.document_processor import processor
    def fail(file, file_type=None):
        raise RuntimeError("invalid PDF")
    monkeypatch.setattr(processor, "process", fail)
    with pytest.raises(grpc.RpcError) as error:
        prepare(document_rpc, session_id, files=[pb.UploadedFile(filename="bad.pdf", content=b"pdf")])
    assert error.value.code() == grpc.StatusCode.INTERNAL
    assert "bad.pdf" in error.value.details() and "invalid PDF" in error.value.details()
    assert resources[1] == []


def test_failed_upload_does_not_poison_session(resources, document_rpc, s3_store, session_id, monkeypatch):
    """解析失败不写 raw：同一会话的后续上传不被坏文件永久卡死。"""
    from document_service.document_processor import processor
    from document_service.document_processor.schemas import ProcessResult
    def parse(file, file_type=None):
        if file.filename == "bad.pdf":
            raise RuntimeError("corrupt document")
        return ProcessResult(filename=file.filename, html="<p>ok</p>")
    monkeypatch.setattr(processor, "process", parse)
    with pytest.raises(grpc.RpcError) as error:
        prepare(document_rpc, session_id, files=[pb.UploadedFile(filename="bad.pdf", content=b"corrupt")])
    assert error.value.code() == grpc.StatusCode.INTERNAL
    bucket = f"res_{session_id}"
    assert s3_store.list_objects(bucket, prefix="raw/") == []
    result = prepare(document_rpc, session_id, files=[pb.UploadedFile(filename="a.docx", content=_docx_bytes())])
    refs = list(result.resource_path)
    assert [ref.location.rsplit("/", 1)[1] for ref in refs if ref.type == "raw"] == ["a.docx"]


@pytest.mark.asyncio
async def test_server_warms_embedding_before_becoming_available(monkeypatch):
    from types import SimpleNamespace
    from document_service.main import create_server
    calls = []
    monkeypatch.setattr(embedding_model, "get_embedder", lambda: SimpleNamespace(
        encode=lambda texts: calls.append(texts)))
    server = await create_server(warmup=True)
    try:
        assert len(calls) == 1 and len(calls[0]) == 1
    finally:
        await server.stop(0)
